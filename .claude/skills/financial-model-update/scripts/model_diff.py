#!/usr/bin/env python3
"""Render a read-only before->after diff for a proposed model change set.

This script NEVER writes the workbook. It opens it read-only, looks up the
current value of each targeted cell/named-range, and prints a diff table so
the analyst can approve or reject before any write happens.

Usage:
    python3 model_diff.py <model.xlsx> --proposed proposed.json

proposed.json schema:
    [{"target": "Assumptions!B12" | "FY26_Rev",
      "new": 11650,
      "source": "Q1'26 PR ... [retrieved 2026-05-16T14:02Z]"}]

Requires openpyxl. If it's not installed, exits non-zero with guidance so
the skill can fall back to a manual analyst-built diff (never a silent skip).
"""
import sys
import json
import argparse

try:
    from openpyxl import load_workbook
    from openpyxl.utils.cell import coordinate_to_tuple  # noqa: F401 (validation)
except ImportError:
    print("ERROR: openpyxl not installed. Run `pip install openpyxl`, or "
          "build the before->after diff manually with the analyst. "
          "Do NOT proceed to write the model without a reviewed diff.",
          file=sys.stderr)
    sys.exit(3)


def resolve(wb, target):
    """Return (sheet, coord, current_value) for 'Sheet!A1' or a named range."""
    if "!" in target:
        sheet_name, coord = target.split("!", 1)
        ws = wb[sheet_name]
        return sheet_name, coord, ws[coord].value
    dn = wb.defined_names.get(target)
    if dn is None:
        return None, target, None
    for sheet_name, coord in dn.destinations:
        return sheet_name, coord, wb[sheet_name][coord].value
    return None, target, None


def main():
    p = argparse.ArgumentParser()
    p.add_argument("model")
    p.add_argument("--proposed", required=True)
    args = p.parse_args()

    with open(args.proposed) as fh:
        changes = json.load(fh)

    wb = load_workbook(args.model, data_only=False, read_only=True)

    rows, problems = [], []
    for ch in changes:
        target = ch["target"]
        sheet, coord, old = resolve(wb, target)
        if sheet is None:
            problems.append(f"Could not resolve target: {target}")
            continue
        old_v, new_v = old, ch["new"]
        flag = ""
        try:
            if old_v not in (None, 0) and isinstance(old_v, (int, float)):
                if abs((float(new_v) - float(old_v)) / float(old_v)) > 0.25:
                    flag = "  <-- LARGE CHANGE (>25%)"
        except (TypeError, ValueError):
            pass
        rows.append((f"{sheet}!{coord}", old_v, new_v, ch.get("source", "MISSING"), flag))

    print("\nProposed model changes (READ-ONLY — nothing written):\n")
    print(f"{'Cell':<22}{'Old':<14}{'New':<14}Source")
    print("-" * 80)
    for cell, old_v, new_v, src, flag in rows:
        print(f"{cell:<22}{str(old_v):<14}{str(new_v):<14}{src}{flag}")
        if src == "MISSING":
            problems.append(f"No source provided for {cell} — block this change.")

    if problems:
        print("\nPROBLEMS — resolve before requesting approval:")
        for pr in problems:
            print(f"  - {pr}")

    print("\nNext: present this to the analyst. Do NOT write the workbook "
          "until they explicitly approve the change set.")


if __name__ == "__main__":
    main()
